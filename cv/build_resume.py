#!/usr/bin/env python3
"""Build jason-herring-resume.pdf and .docx from resume.json.

  python3 cv/build_resume.py

The PDF is printed from resume.html.j2 via headless Chrome; the .docx is a
single-column, table-free ATS layout built with python-docx. Both read the same
content file, so an update is a resume.json edit plus a re-run.

Requires: jinja2, python-docx, and Google Chrome.
"""

import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Pt, RGBColor
from jinja2 import Environment, FileSystemLoader, select_autoescape
from markupsafe import Markup, escape

CV_DIR = Path(__file__).resolve().parent
CONTENT = CV_DIR / "resume.json"
TEMPLATE = "resume.html.j2"
HTML_OUT = CV_DIR / ".resume.build.html"
PDF_OUT = CV_DIR / "jason-herring-resume.pdf"
DOCX_OUT = CV_DIR / "jason-herring-resume.docx"

CHROME_CANDIDATES = [
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
    "google-chrome",
    "chromium",
]

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")

# Word measurements: page 8.5in wide, 0.6in margins -> 7.3in of text.
DOCX_RIGHT_TAB = Pt(7.3 * 72)
GREY = RGBColor(0x6A, 0x6A, 0x6A)
DIM = RGBColor(0x5C, 0x5C, 0x5C)


# --------------------------------------------------------------------------
# shared: **bold** inline markup
# --------------------------------------------------------------------------
def split_bold(text):
    """Yield (chunk, is_bold) pairs, splitting on **bold** spans."""
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            yield text[pos : m.start()], False
        yield m.group(1), True
        pos = m.end()
    if pos < len(text):
        yield text[pos:], False


def md_filter(text):
    """Jinja filter: escape, then render **bold** as <strong>."""
    return Markup(
        "".join(
            f"<strong>{escape(chunk)}</strong>" if bold else str(escape(chunk))
            for chunk, bold in split_bold(text)
        )
    )


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def find_chrome():
    for candidate in CHROME_CANDIDATES:
        if Path(candidate).exists():
            return candidate
        found = shutil.which(candidate)
        if found:
            return found
    sys.exit("Could not find Google Chrome or Chromium — needed to print the PDF.")


def build_pdf(data):
    env = Environment(
        loader=FileSystemLoader(CV_DIR),
        autoescape=select_autoescape(["html", "j2"]),
    )
    env.filters["md"] = md_filter
    HTML_OUT.write_text(env.get_template(TEMPLATE).render(**data), encoding="utf-8")

    subprocess.run(
        [
            find_chrome(),
            "--headless",
            "--disable-gpu",
            "--no-sandbox",
            "--no-pdf-header-footer",
            f"--print-to-pdf={PDF_OUT}",
            HTML_OUT.as_uri(),
        ],
        check=True,
        capture_output=True,
    )
    HTML_OUT.unlink()
    print(f"wrote {PDF_OUT.relative_to(CV_DIR.parent)}")


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------
def para(doc, space_before=0, space_after=0, left_indent=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.08
    if left_indent is not None:
        pf.left_indent = Pt(left_indent)
    return p


def run(p, text, size=9.5, bold=False, italic=False, color=None):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r


def rich(p, text, size=9.5, color=None):
    """Add text to a paragraph, honouring **bold** spans."""
    for chunk, bold in split_bold(text):
        run(p, chunk, size=size, bold=bold, color=color)


def section_heading(doc, text):
    p = para(doc, space_before=11, space_after=3)
    run(p, text.upper(), size=8.5, bold=True)


def bullet(doc, text, size=9.5, indent=14):
    p = para(doc, space_after=2, left_indent=indent)
    run(p, "•  ", size=size, color=GREY)
    rich(p, text, size=size)
    return p


def build_docx(data):
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.5)

    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(0.5 * 72)
        section.left_margin = section.right_margin = Pt(0.6 * 72)

    # header
    p = para(doc)
    run(p, data["name"], size=22, bold=True)
    p = para(doc, space_after=1)
    run(p, data["title"], size=11)
    p = para(doc, space_after=4)
    run(p, " · ".join(data["contact"]), size=8.5, color=GREY)

    section_heading(doc, "Summary")
    rich(para(doc), data["summary"])

    section_heading(doc, "Core Skills")
    para(doc).add_run(" · ".join(data["skills"])).font.size = Pt(9.5)

    section_heading(doc, "Experience")
    for job in data["experience"]:
        p = para(doc, space_before=8)
        p.paragraph_format.tab_stops.add_tab_stop(
            DOCX_RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT
        )
        run(p, f"{job['company']} — {job['role']}", size=10.5, bold=True)
        run(p, "\t")
        run(p, job["dates"], size=8.5, color=GREY)

        if job.get("context"):
            p = para(doc, space_after=2)
            run(p, job["context"], size=8.5, italic=True, color=DIM)

        if job.get("shipped"):
            p = para(doc, space_before=5, space_after=1)
            run(p, job["shipped_label"].upper(), size=8, bold=True, color=GREY)
            for s in job["shipped"]:
                p = para(doc, space_after=2, left_indent=14)
                run(p, "•  ", color=GREY)
                run(p, s["product"], bold=True)
                if s.get("note"):
                    run(p, f" ({s['note']})", color=GREY)
                run(p, " — ")
                rich(p, s["text"])

        for i, b in enumerate(job.get("bullets", [])):
            p = bullet(doc, b)
            # keep the plain bullets from reading as more shipped-work entries
            if i == 0 and job.get("shipped"):
                p.paragraph_format.space_before = Pt(6)

    section_heading(doc, "Education")
    for e in data["education"]:
        run(para(doc, space_after=1), e)

    section_heading(doc, "Awards & Recognition")
    run(para(doc), data["awards"])

    doc.save(DOCX_OUT)
    print(f"wrote {DOCX_OUT.relative_to(CV_DIR.parent)}")


def main():
    data = json.loads(CONTENT.read_text(encoding="utf-8"))
    build_pdf(data)
    build_docx(data)


if __name__ == "__main__":
    main()
